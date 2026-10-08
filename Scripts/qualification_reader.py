import json
import re
import warnings
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from line_schedule_reader import read_line_schedule


PROJECT_ROOT = Path(__file__).resolve().parent.parent
SETTINGS_PATH = PROJECT_ROOT / "Config" / "settings.json"

EXPECTED_QUALIFICATIONS = {
    "DUMOD0",
    "DUMOD1",
    "CUMOD0",
    "CUMOD1",
    "CUMOD2",
    "CUMOD3",
    "UCMOD0",
    "UCMOD1",
    "UCMOD2",
    "UCMOD3",
    "UCMOD4",
}

REQUIRED_HEADERS = {
    "Labor",
    "Name",
    "Effective Date",
    "Status",
}

warnings.filterwarnings(
    "ignore",
    message="Workbook contains no default style",
)


def load_settings():
    if not SETTINGS_PATH.exists():
        raise FileNotFoundError(
            f"Settings file not found: {SETTINGS_PATH}"
        )

    with SETTINGS_PATH.open(
        mode="r",
        encoding="utf-8",
    ) as settings_file:
        return json.load(settings_file)


def get_qualification_folder(settings):
    qualification_folder = Path(
        settings["qualification_folder"]
    )

    if not qualification_folder.exists():
        raise FileNotFoundError(
            "Qualification folder not found:\n"
            f"{qualification_folder}"
        )

    if not qualification_folder.is_dir():
        raise NotADirectoryError(
            "Qualification path is not a folder:\n"
            f"{qualification_folder}"
        )

    return qualification_folder


def find_qualification_files(qualification_folder):
    supported_extensions = {
        ".xlsx",
        ".xlsm",
    }

    return sorted(
        file
        for file in qualification_folder.iterdir()
        if file.is_file()
        and file.suffix.lower() in supported_extensions
        and not file.name.startswith("~$")
    )


def identify_qualification_code(file_path):
    filename_upper = file_path.stem.upper()

    matches = [
        qualification_code
        for qualification_code in EXPECTED_QUALIFICATIONS
        if qualification_code in filename_upper
    ]

    if len(matches) == 1:
        return matches[0]

    if not matches:
        return "UNRECOGNIZED"

    return "MULTIPLE MATCHES"


def clean_sso(value):
    if value is None:
        return ""

    if isinstance(value, int):
        sso = str(value)
    elif isinstance(value, float) and value.is_integer():
        sso = str(int(value))
    else:
        sso = str(value).strip()
        sso = re.sub(r"\.0$", "", sso)

    return sso if re.fullmatch(r"\d{9}", sso) else ""


def clean_text(value):
    if value is None:
        return ""

    return re.sub(r"\s+", " ", str(value).strip())


def parse_date(value):
    if value in (None, ""):
        return None

    if isinstance(value, datetime):
        return value.date()

    if hasattr(value, "year") and hasattr(value, "month"):
        return value

    text_value = str(value).strip()

    date_formats = (
        "%m/%d/%y",
        "%m/%d/%Y",
        "%Y-%m-%d",
    )

    for date_format in date_formats:
        try:
            return datetime.strptime(
                text_value,
                date_format,
            ).date()
        except ValueError:
            continue

    return None


def get_header_map(worksheet):
    first_row = next(
        worksheet.iter_rows(
            min_row=1,
            max_row=1,
            values_only=True,
        )
    )

    header_map = {
        clean_text(header): column_index
        for column_index, header in enumerate(first_row)
        if clean_text(header)
    }

    missing_headers = REQUIRED_HEADERS - header_map.keys()

    if missing_headers:
        missing_list = ", ".join(sorted(missing_headers))
        raise ValueError(
            f"Missing required headers in "
            f"worksheet '{worksheet.title}': "
            f"{missing_list}"
        )

    return header_map


def read_qualification_file(
    file_path,
    qualification_code,
    roster_ssos,
):
    workbook = load_workbook(
        file_path,
        read_only=True,
        data_only=True,
    )

    if "List of Labor" not in workbook.sheetnames:
        workbook.close()
        raise ValueError(
            f"Worksheet 'List of Labor' was not found in "
            f"{file_path.name}"
        )

    worksheet = workbook["List of Labor"]
    worksheet.reset_dimensions()

    header_map = get_header_map(worksheet)

    matched_records = []
    inactive_records = []
    invalid_date_records = []

    for row_number, row in enumerate(
        worksheet.iter_rows(
            min_row=2,
            values_only=True,
        ),
        start=2,
    ):
        labor_value = row[header_map["Labor"]]
        name_value = row[header_map["Name"]]
        date_value = row[header_map["Effective Date"]]
        status_value = row[header_map["Status"]]

        sso = clean_sso(labor_value)

        if not sso or sso not in roster_ssos:
            continue

        qualification_name = clean_text(name_value)
        status = clean_text(status_value).upper()
        effective_date = parse_date(date_value)

        record = {
            "sso": sso,
            "qualification_name": qualification_name,
            "qualification": qualification_code,
            "effective_date": effective_date,
            "status": status,
            "source_file": file_path.name,
            "source_row": row_number,
        }

        if status != "ACTIVE":
            inactive_records.append(record)
            continue

        if effective_date is None:
            invalid_date_records.append(record)

        matched_records.append(record)

    workbook.close()

    return (
        matched_records,
        inactive_records,
        invalid_date_records,
    )


def select_current_records(records):
    selected_records = {}
    duplicate_records = []

    for record in records:
        key = (
            record["sso"],
            record["qualification"],
        )

        if key not in selected_records:
            selected_records[key] = record
            continue

        current_record = selected_records[key]

        duplicate_records.append(
            {
                "sso": record["sso"],
                "qualification": record["qualification"],
                "kept_date": current_record["effective_date"],
                "additional_date": record["effective_date"],
                "source_file": record["source_file"],
                "source_row": record["source_row"],
            }
        )

        current_date = current_record["effective_date"]
        new_date = record["effective_date"]

        if current_date is None and new_date is not None:
            selected_records[key] = record
        elif (
            current_date is not None
            and new_date is not None
            and new_date > current_date
        ):
            selected_records[key] = record

    return selected_records, duplicate_records


def read_all_qualifications():
    technicians, line_schedule_exceptions = (
        read_line_schedule()
    )

    roster_by_sso = {
        technician["sso"]: technician
        for technician in technicians
    }

    roster_ssos = set(roster_by_sso)

    settings = load_settings()
    qualification_folder = get_qualification_folder(
        settings
    )
    qualification_files = find_qualification_files(
        qualification_folder
    )

    file_by_code = {}
    file_exceptions = []

    for file_path in qualification_files:
        qualification_code = identify_qualification_code(
            file_path
        )

        if qualification_code not in EXPECTED_QUALIFICATIONS:
            file_exceptions.append(
                {
                    "file": file_path.name,
                    "issue": (
                        "Qualification code could not be "
                        "identified from filename"
                    ),
                }
            )
            continue

        if qualification_code in file_by_code:
            file_exceptions.append(
                {
                    "file": file_path.name,
                    "issue": (
                        "Duplicate qualification file for "
                        f"{qualification_code}"
                    ),
                }
            )
            continue

        file_by_code[qualification_code] = file_path

    missing_codes = (
        EXPECTED_QUALIFICATIONS - file_by_code.keys()
    )

    if missing_codes:
        missing_list = ", ".join(sorted(missing_codes))
        raise FileNotFoundError(
            "Missing qualification files: "
            f"{missing_list}"
        )

    all_active_records = []
    inactive_records = []
    invalid_date_records = []

    for qualification_code in sorted(
        EXPECTED_QUALIFICATIONS
    ):
        file_path = file_by_code[qualification_code]

        (
            matched,
            inactive,
            invalid_dates,
        ) = read_qualification_file(
            file_path=file_path,
            qualification_code=qualification_code,
            roster_ssos=roster_ssos,
        )

        all_active_records.extend(matched)
        inactive_records.extend(inactive)
        invalid_date_records.extend(invalid_dates)

    selected_records, duplicate_records = (
        select_current_records(all_active_records)
    )

    return {
        "technicians": technicians,
        "roster_by_sso": roster_by_sso,
        "qualification_records": selected_records,
        "line_schedule_exceptions": (
            line_schedule_exceptions
        ),
        "file_exceptions": file_exceptions,
        "inactive_records": inactive_records,
        "invalid_date_records": invalid_date_records,
        "duplicate_records": duplicate_records,
    }


def display_match_summary():
    results = read_all_qualifications()

    technicians = results["technicians"]
    records = results["qualification_records"]

    print(f"Roster technicians: {len(technicians)}")
    print(
        "Active roster qualification records: "
        f"{len(records)}"
    )
    print()

    for qualification_code in sorted(
        EXPECTED_QUALIFICATIONS
    ):
        matched_ssos = {
            sso
            for sso, qualification
            in records
            if qualification == qualification_code
        }

        print(
            f"{qualification_code}: "
            f"{len(matched_ssos)} of "
            f"{len(technicians)} roster technicians"
        )

    qualified_ssos = {
        sso
        for sso, qualification in records
    }

    no_qualification_records = [
        technician
        for technician in technicians
        if technician["sso"] not in qualified_ssos
    ]

    print()
    print(
        "Technicians with no active qualification "
        f"records: {len(no_qualification_records)}"
    )

    for technician in no_qualification_records:
        print(
            f"  {technician['sso']} | "
            f"{technician['name']} | "
            f"{technician['schedule_group']}"
        )

    print()
    print(
        "Inactive matched records: "
        f"{len(results['inactive_records'])}"
    )
    print(
        "Invalid effective dates: "
        f"{len(results['invalid_date_records'])}"
    )
    print(
        "Duplicate active records: "
        f"{len(results['duplicate_records'])}"
    )
    print(
        "Line Schedule exceptions: "
        f"{len(results['line_schedule_exceptions'])}"
    )
    print(
        "Qualification file exceptions: "
        f"{len(results['file_exceptions'])}"
    )


if __name__ == "__main__":
    display_match_summary()