import json
import shutil
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)
from openpyxl.utils import get_column_letter

from coverage_processor import (
    VENUE_MODULES,
    calculate_dashboard_results,
)
from qualification_reader import read_all_qualifications


PROJECT_ROOT = Path(__file__).resolve().parent.parent
SETTINGS_PATH = PROJECT_ROOT / "Config" / "settings.json"

ALL_QUALIFICATIONS = [
    "DUMOD0",
    "DUMOD1",
    "CUMOD0",
    "CUMOD1",
    "CUMOD2",
    "CUMOD3",
    "UCMOD0",
    "UCMOD1",
    "UCMOD2",
    "UCMOD3",
    "UCMOD4",
]

DARK_BLUE = "1F4E78"
MEDIUM_BLUE = "5B9BD5"
LIGHT_BLUE = "D9EAF7"
GREEN = "70AD47"
LIGHT_GREEN = "C6E0B4"
YELLOW = "FFE699"
ORANGE = "F4B183"
LIGHT_RED = "F4CCCC"
WHITE = "FFFFFF"
BLACK = "000000"
GRAY = "D9E1F2"
LIGHT_GRAY = "E7E6E6"

THIN_GRAY = Side(
    style="thin",
    color="B7B7B7",
)

HEADER_FILL = PatternFill(
    fill_type="solid",
    fgColor=DARK_BLUE,
)

SUBHEADER_FILL = PatternFill(
    fill_type="solid",
    fgColor=MEDIUM_BLUE,
)

LINKED_FILL = PatternFill(
    fill_type="solid",
    fgColor=LIGHT_GREEN,
)

STATIC_FILL = PatternFill(
    fill_type="solid",
    fgColor=LIGHT_GRAY,
)

CAUTION_FILL = PatternFill(
    fill_type="solid",
    fgColor=YELLOW,
)

ERROR_FILL = PatternFill(
    fill_type="solid",
    fgColor=LIGHT_RED,
)


def load_settings():
    with SETTINGS_PATH.open(
        mode="r",
        encoding="utf-8",
    ) as settings_file:
        return json.load(settings_file)


def get_output_folder(settings):
    output_folder = Path(settings["output_folder"])

    if not output_folder.is_absolute():
        output_folder = PROJECT_ROOT / output_folder

    output_folder.mkdir(
        parents=True,
        exist_ok=True,
    )

    return output_folder

def archive_existing_outputs(
    output_folder,
    archive_folder,
):
    existing_workbooks = [
        file_path
        for file_path in output_folder.glob("*.xlsx")
        if not file_path.name.startswith("~$")
    ]

    if not existing_workbooks:
        return None, 0

    archive_timestamp = datetime.now().strftime(
        "%Y-%m-%d_%H-%M-%S"
    )

    timestamped_archive_folder = (
        archive_folder / archive_timestamp
    )

    timestamped_archive_folder.mkdir(
        parents=True,
        exist_ok=True,
    )

    archived_count = 0

    for file_path in existing_workbooks:
        destination = (
            timestamped_archive_folder
            / file_path.name
        )

        shutil.move(
            str(file_path),
            str(destination),
        )

        archived_count += 1

    return timestamped_archive_folder, archived_count


def apply_title_style(cell):
    cell.fill = HEADER_FILL
    cell.font = Font(
        color=WHITE,
        bold=True,
        size=16,
    )
    cell.alignment = Alignment(
        horizontal="left",
        vertical="center",
    )


def apply_header_style(cell):
    cell.fill = HEADER_FILL
    cell.font = Font(
        color=WHITE,
        bold=True,
    )
    cell.alignment = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=True,
    )
    cell.border = Border(
        bottom=THIN_GRAY,
    )


def apply_subheader_style(cell):
    cell.fill = SUBHEADER_FILL
    cell.font = Font(
        color=WHITE,
        bold=True,
    )
    cell.alignment = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=True,
    )


def apply_section_style(
    worksheet,
    row_number,
    start_column,
    end_column,
):
    for column_number in range(
        start_column,
        end_column + 1,
    ):
        cell = worksheet.cell(
            row=row_number,
            column=column_number,
        )
        cell.fill = HEADER_FILL
        cell.font = Font(
            color=WHITE,
            bold=True,
        )
        cell.alignment = Alignment(
            horizontal="left",
            vertical="center",
        )


def set_column_widths(worksheet, widths):
    for column_letter, width in widths.items():
        worksheet.column_dimensions[
            column_letter
        ].width = width


def find_module_result(venue_result, qualification):
    return next(
        (
            module_result
            for module_result
            in venue_result["module_results"]
            if module_result["qualification"]
            == qualification
        ),
        None,
    )


def write_percentage_cell(cell, value):
    cell.value = value
    cell.number_format = "0%"
    cell.alignment = Alignment(
        horizontal="center",
        vertical="center",
    )
    cell.font = Font(
        bold=True,
        color=BLACK,
    )

    if value >= 0.90:
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor=GREEN,
        )
    elif value >= 0.80:
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor=LIGHT_GREEN,
        )
    elif value >= 0.70:
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor=YELLOW,
        )
    elif value >= 0.60:
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor=ORANGE,
        )
    else:
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor=LIGHT_RED,
        )


def add_percentage_colors(
    worksheet,
    cell_range,
):
    worksheet.conditional_formatting.add(
        cell_range,
        CellIsRule(
            operator="greaterThanOrEqual",
            formula=["0.9"],
            fill=PatternFill(
                fill_type="solid",
                fgColor=GREEN,
            ),
        ),
    )

    worksheet.conditional_formatting.add(
        cell_range,
        CellIsRule(
            operator="between",
            formula=["0.8", "0.899999"],
            fill=PatternFill(
                fill_type="solid",
                fgColor=LIGHT_GREEN,
            ),
        ),
    )

    worksheet.conditional_formatting.add(
        cell_range,
        CellIsRule(
            operator="between",
            formula=["0.7", "0.799999"],
            fill=PatternFill(
                fill_type="solid",
                fgColor=YELLOW,
            ),
        ),
    )

    worksheet.conditional_formatting.add(
        cell_range,
        CellIsRule(
            operator="between",
            formula=["0.6", "0.699999"],
            fill=PatternFill(
                fill_type="solid",
                fgColor=ORANGE,
            ),
        ),
    )

    worksheet.conditional_formatting.add(
        cell_range,
        CellIsRule(
            operator="lessThan",
            formula=["0.6"],
            fill=PatternFill(
                fill_type="solid",
                fgColor=LIGHT_RED,
            ),
        ),
    )


def create_dashboard(
    workbook,
    dashboard_results,
):
    worksheet = workbook.active
    worksheet.title = "Dashboard"
    worksheet.sheet_view.showGridLines = False

    worksheet.merge_cells("A1:H1")
    worksheet["A1"] = (
        "Dark Universe Qualification Coverage"
    )
    apply_title_style(worksheet["A1"])
    worksheet.row_dimensions[1].height = 26

    worksheet["A2"] = "Report generated"
    worksheet["B2"] = datetime.now()
    worksheet["B2"].number_format = (
        "mm/dd/yyyy h:mm AM/PM"
    )

    worksheet["A3"] = "Valid technicians"
    worksheet["B3"] = sum(
        section["population_count"]
        for section
        in dashboard_results["sections"].values()
    )

    headers = [
        "Venue",
        "MOD0",
        "MOD1",
        "MOD2",
        "MOD3",
        "MOD4",
        "Total",
        "Population",
    ]

    current_row = 5
    percentage_cells = []

    dashboard_sections = [
        (
            "3rd Shift",
            [
                (
                    "Unchained",
                    dashboard_results["sections"][
                        "3rd Shift Unchained"
                    ]["venues"]["Unchained"],
                ),
                (
                    "Curse of the Werewolf",
                    dashboard_results["sections"][
                        "3rd Shift Wolf"
                    ]["venues"]["Wolf"],
                ),
                (
                    "Darkmoor",
                    dashboard_results["sections"][
                        "3rd Shift Wolf"
                    ]["venues"]["Darkmoor"],
                ),
            ],
        ),
        (
            "Operational",
            [
                (
                    "Unchained",
                    dashboard_results["sections"][
                        "Operational"
                    ]["venues"]["Unchained"],
                ),
                (
                    "Curse of the Werewolf",
                    dashboard_results["sections"][
                        "Operational"
                    ]["venues"]["Wolf"],
                ),
                (
                    "Darkmoor",
                    dashboard_results["sections"][
                        "Operational"
                    ]["venues"]["Darkmoor"],
                ),
            ],
        ),
        (
            "Total",
            [
                (
                    "Unchained",
                    dashboard_results["totals"][
                        "Unchained"
                    ],
                ),
                (
                    "Curse of the Werewolf",
                    dashboard_results["totals"][
                        "Wolf"
                    ],
                ),
                (
                    "Darkmoor",
                    dashboard_results["totals"][
                        "Darkmoor"
                    ],
                ),
            ],
        ),
    ]

    for section_name, venue_rows in dashboard_sections:
        worksheet.merge_cells(
            start_row=current_row,
            start_column=1,
            end_row=current_row,
            end_column=8,
        )
        worksheet.cell(
            row=current_row,
            column=1,
            value=section_name,
        )
        apply_section_style(
            worksheet,
            current_row,
            1,
            8,
        )
        current_row += 1

        for column_number, header in enumerate(
            headers,
            start=1,
        ):
            cell = worksheet.cell(
                row=current_row,
                column=column_number,
                value=header,
            )
            apply_subheader_style(cell)

        current_row += 1

        for venue_name, venue_result in venue_rows:
            worksheet.cell(
                row=current_row,
                column=1,
                value=venue_name,
            ).font = Font(bold=True)

            prefix = {
                "Unchained": "UCMOD",
                "Curse of the Werewolf": "CUMOD",
                "Darkmoor": "DUMOD",
            }[venue_name]

            for module_number in range(5):
                qualification = (
                    f"{prefix}{module_number}"
                )

                module_result = find_module_result(
                    venue_result,
                    qualification,
                )

                if module_result is None:
                    continue

                cell = worksheet.cell(
                    row=current_row,
                    column=module_number + 2,
                )

                write_percentage_cell(
                    cell,
                    module_result["percentage"],
                )

                percentage_cells.append(
                    cell.coordinate
                )

                cell.comment = None

            total_cell = worksheet.cell(
                row=current_row,
                column=7,
            )
            write_percentage_cell(
                total_cell,
                venue_result["total_percentage"],
            )
            percentage_cells.append(
                total_cell.coordinate
            )

            worksheet.cell(
                row=current_row,
                column=8,
                value=venue_result[
                    "population_count"
                ],
            ).alignment = Alignment(
                horizontal="center"
            )

            current_row += 1

        current_row += 1

    for cell_coordinate in percentage_cells:
        add_percentage_colors(
            worksheet,
            cell_coordinate,
        )

    set_column_widths(
        worksheet,
        {
            "A": 27,
            "B": 11,
            "C": 11,
            "D": 11,
            "E": 11,
            "F": 11,
            "G": 11,
            "H": 12,
        },
    )

    worksheet.freeze_panes = "A5"
    worksheet.auto_filter.ref = (
        f"A6:H{current_row - 1}"
    )


def calculate_highest_level(
    sso,
    qualification_records,
    modules,
):
    achieved_modules = [
        qualification
        for qualification in modules
        if (sso, qualification)
        in qualification_records
    ]

    if not achieved_modules:
        return "None"

    highest_contiguous_level = -1

    for level, qualification in enumerate(modules):
        if qualification in achieved_modules:
            highest_contiguous_level = level
        else:
            break

    if highest_contiguous_level == -1:
        return "None"

    return f"Through Mod {highest_contiguous_level}"


def find_qualification_gaps(
    sso,
    qualification_records,
):
    gaps = []

    for venue, modules in VENUE_MODULES.items():
        missing_lower_module = None

        for qualification in modules:
            has_qualification = (
                sso,
                qualification,
            ) in qualification_records

            if not has_qualification:
                if missing_lower_module is None:
                    missing_lower_module = qualification
                continue

            if missing_lower_module is not None:
                gaps.append(
                    f"{venue}: has {qualification}, "
                    f"missing {missing_lower_module}"
                )

    return "; ".join(gaps)


def create_technician_sheet(
    workbook,
    qualification_data,
):
    worksheet = workbook.create_sheet(
        "Technician Qualifications"
    )
    worksheet.sheet_view.showGridLines = False

    headers = [
        "Name",
        "SSO",
        "Schedule Group",
    ]

    for qualification in ALL_QUALIFICATIONS:
        headers.extend(
            [
                qualification,
                f"{qualification} Date",
            ]
        )

    headers.extend(
        [
            "Darkmoor Level",
            "Wolf Level",
            "Unchained Level",
            "Review Flag",
        ]
    )

    for column_number, header in enumerate(
        headers,
        start=1,
    ):
        cell = worksheet.cell(
            row=1,
            column=column_number,
            value=header,
        )
        apply_header_style(cell)

    qualification_records = qualification_data[
        "qualification_records"
    ]

    technicians = sorted(
        qualification_data["technicians"],
        key=lambda technician: (
            technician["schedule_group"],
            technician["name"],
        ),
    )

    for row_number, technician in enumerate(
        technicians,
        start=2,
    ):
        sso = technician["sso"]

        row_values = [
            technician["name"],
            sso,
            technician["schedule_group"],
        ]

        for qualification in ALL_QUALIFICATIONS:
            record = qualification_records.get(
                (
                    sso,
                    qualification,
                )
            )

            if record:
                row_values.extend(
                    [
                        "Yes",
                        record["effective_date"],
                    ]
                )
            else:
                row_values.extend(
                    [
                        "",
                        "",
                    ]
                )

        row_values.extend(
            [
                calculate_highest_level(
                    sso,
                    qualification_records,
                    VENUE_MODULES["Darkmoor"],
                ),
                calculate_highest_level(
                    sso,
                    qualification_records,
                    VENUE_MODULES["Wolf"],
                ),
                calculate_highest_level(
                    sso,
                    qualification_records,
                    VENUE_MODULES["Unchained"],
                ),
                find_qualification_gaps(
                    sso,
                    qualification_records,
                ),
            ]
        )

        for column_number, value in enumerate(
            row_values,
            start=1,
        ):
            cell = worksheet.cell(
                row=row_number,
                column=column_number,
                value=value,
            )

            if isinstance(value, datetime):
                cell.number_format = "mm/dd/yyyy"
            elif hasattr(value, "year"):
                cell.number_format = "mm/dd/yyyy"

    for row_number in range(
        2,
        worksheet.max_row + 1,
    ):
        for column_number in range(
            4,
            4 + len(ALL_QUALIFICATIONS) * 2,
            2,
        ):
            cell = worksheet.cell(
                row=row_number,
                column=column_number,
            )

            if cell.value == "Yes":
                cell.fill = LINKED_FILL
                cell.font = Font(
                    color="008000",
                    bold=True,
                )
                cell.alignment = Alignment(
                    horizontal="center",
                )

        review_cell = worksheet.cell(
            row=row_number,
            column=worksheet.max_column,
        )

        if review_cell.value:
            review_cell.fill = ERROR_FILL

    worksheet.freeze_panes = "D2"
    worksheet.auto_filter.ref = (
        f"A1:{get_column_letter(worksheet.max_column)}"
        f"{worksheet.max_row}"
    )

    worksheet.column_dimensions["A"].width = 26
    worksheet.column_dimensions["B"].width = 13
    worksheet.column_dimensions["C"].width = 27

    for column_number in range(
        4,
        worksheet.max_column + 1,
    ):
        column_letter = get_column_letter(
            column_number
        )

        if "Date" in str(
            worksheet.cell(
                row=1,
                column=column_number,
            ).value
        ):
            worksheet.column_dimensions[
                column_letter
            ].width = 14
        else:
            worksheet.column_dimensions[
                column_letter
            ].width = 16


def create_venue_summary(
    workbook,
    dashboard_results,
):
    worksheet = workbook.create_sheet(
        "Venue Summary"
    )
    worksheet.sheet_view.showGridLines = False

    headers = [
        "Section",
        "Venue",
        "Qualification",
        "Qualified",
        "Population",
        "Percentage",
    ]

    for column_number, header in enumerate(
        headers,
        start=1,
    ):
        apply_header_style(
            worksheet.cell(
                row=1,
                column=column_number,
                value=header,
            )
        )

    current_row = 2

    sections = []

    for section_name, section_result in (
        dashboard_results["sections"].items()
    ):
        for venue_name, venue_result in (
            section_result["venues"].items()
        ):
            sections.append(
                (
                    section_name,
                    venue_name,
                    venue_result,
                )
            )

    for venue_name, venue_result in (
        dashboard_results["totals"].items()
    ):
        sections.append(
            (
                "Total",
                venue_name,
                venue_result,
            )
        )

    for section_name, venue_name, venue_result in sections:
        for module_result in venue_result[
            "module_results"
        ]:
            worksheet.cell(
                row=current_row,
                column=1,
                value=section_name,
            )
            worksheet.cell(
                row=current_row,
                column=2,
                value=venue_name,
            )
            worksheet.cell(
                row=current_row,
                column=3,
                value=module_result[
                    "qualification"
                ],
            )
            worksheet.cell(
                row=current_row,
                column=4,
                value=module_result[
                    "qualified_count"
                ],
            )
            worksheet.cell(
                row=current_row,
                column=5,
                value=module_result[
                    "population_count"
                ],
            )

            percentage_cell = worksheet.cell(
                row=current_row,
                column=6,
                value=module_result[
                    "percentage"
                ],
            )
            percentage_cell.number_format = "0.0%"

            current_row += 1

    add_percentage_colors(
        worksheet,
        f"F2:F{worksheet.max_row}",
    )

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = (
        f"A1:F{worksheet.max_row}"
    )

    set_column_widths(
        worksheet,
        {
            "A": 24,
            "B": 18,
            "C": 16,
            "D": 13,
            "E": 13,
            "F": 14,
        },
    )


def create_exceptions_sheet(
    workbook,
    qualification_data,
):
    worksheet = workbook.create_sheet(
        "Exceptions"
    )
    worksheet.sheet_view.showGridLines = False

    headers = [
        "Exception Type",
        "SSO",
        "Name/File",
        "Qualification",
        "Issue",
        "Source Row",
    ]

    for column_number, header in enumerate(
        headers,
        start=1,
    ):
        apply_header_style(
            worksheet.cell(
                row=1,
                column=column_number,
                value=header,
            )
        )

    exception_rows = []

    for exception in qualification_data[
        "line_schedule_exceptions"
    ]:
        exception_rows.append(
            [
                "Line Schedule",
                "",
                exception.get("name", ""),
                "",
                exception.get("issue", ""),
                exception.get("row", ""),
            ]
        )

    for exception in qualification_data[
        "file_exceptions"
    ]:
        exception_rows.append(
            [
                "Qualification File",
                "",
                exception.get("file", ""),
                "",
                exception.get("issue", ""),
                "",
            ]
        )

    for record in qualification_data[
        "inactive_records"
    ]:
        exception_rows.append(
            [
                "Inactive Qualification",
                record.get("sso", ""),
                record.get(
                    "qualification_name",
                    "",
                ),
                record.get(
                    "qualification",
                    "",
                ),
                f"Status: {record.get('status', '')}",
                record.get("source_row", ""),
            ]
        )

    for record in qualification_data[
        "invalid_date_records"
    ]:
        exception_rows.append(
            [
                "Invalid Date",
                record.get("sso", ""),
                record.get(
                    "qualification_name",
                    "",
                ),
                record.get(
                    "qualification",
                    "",
                ),
                "Invalid or missing effective date",
                record.get("source_row", ""),
            ]
        )

    for record in qualification_data[
        "duplicate_records"
    ]:
        exception_rows.append(
            [
                "Duplicate Qualification",
                record.get("sso", ""),
                "",
                record.get(
                    "qualification",
                    "",
                ),
                "Multiple active records found",
                record.get("source_row", ""),
            ]
        )

    if not exception_rows:
        worksheet.cell(
            row=2,
            column=1,
            value="No exceptions found",
        )
        worksheet.merge_cells(
            start_row=2,
            start_column=1,
            end_row=2,
            end_column=6,
        )
        worksheet["A2"].fill = LINKED_FILL
        worksheet["A2"].font = Font(
            bold=True,
            color="008000",
        )
    else:
        for row_number, row_values in enumerate(
            exception_rows,
            start=2,
        ):
            for column_number, value in enumerate(
                row_values,
                start=1,
            ):
                cell = worksheet.cell(
                    row=row_number,
                    column=column_number,
                    value=value,
                )
                cell.fill = ERROR_FILL

    set_column_widths(
        worksheet,
        {
            "A": 25,
            "B": 13,
            "C": 30,
            "D": 17,
            "E": 40,
            "F": 12,
        },
    )


def create_run_summary(
    workbook,
    qualification_data,
    dashboard_results,
    output_path,
):
    worksheet = workbook.create_sheet(
        "Run Summary"
    )
    worksheet.sheet_view.showGridLines = False

    worksheet.merge_cells("A1:B1")
    worksheet["A1"] = "Qualification Report Run Summary"
    apply_title_style(worksheet["A1"])

    total_exceptions = sum(
        len(exception_list)
        for exception_list
        in dashboard_results["exceptions"].values()
    )

    summary_rows = [
        (
            "Run date and time",
            datetime.now(),
        ),
        (
            "Output workbook",
            output_path.name,
        ),
        (
            "Valid technicians",
            len(qualification_data["technicians"]),
        ),
        (
            "Active qualification matches",
            len(
                qualification_data[
                    "qualification_records"
                ]
            ),
        ),
        (
            "Operational technicians",
            dashboard_results["sections"][
                "Operational"
            ]["population_count"],
        ),
        (
            "3rd Shift Unchained technicians",
            dashboard_results["sections"][
                "3rd Shift Unchained"
            ]["population_count"],
        ),
        (
            "3rd Shift Wolf technicians",
            dashboard_results["sections"][
                "3rd Shift Wolf"
            ]["population_count"],
        ),
        (
            "Unchained applicable population",
            dashboard_results["totals"][
                "Unchained"
            ]["population_count"],
        ),
        (
            "Wolf applicable population",
            dashboard_results["totals"][
                "Wolf"
            ]["population_count"],
        ),
        (
            "Darkmoor applicable population",
            dashboard_results["totals"][
                "Darkmoor"
            ]["population_count"],
        ),
        (
            "Total exceptions",
            total_exceptions,
        ),
    ]

    for row_number, (
        label,
        value,
    ) in enumerate(
        summary_rows,
        start=3,
    ):
        worksheet.cell(
            row=row_number,
            column=1,
            value=label,
        ).font = Font(bold=True)

        value_cell = worksheet.cell(
            row=row_number,
            column=2,
            value=value,
        )

        if isinstance(value, datetime):
            value_cell.number_format = (
                "mm/dd/yyyy h:mm AM/PM"
            )

    set_column_widths(
        worksheet,
        {
            "A": 38,
            "B": 45,
        },
    )


def generate_report():
    settings = load_settings()
    output_folder = get_output_folder(settings)

    archive_folder = Path(
        settings["archive_folder"]
    )

    if not archive_folder.is_absolute():
        archive_folder = (
            PROJECT_ROOT / archive_folder
        )

    archive_folder.mkdir(
        parents=True,
        exist_ok=True,
    )

    archive_path, archived_count = (
        archive_existing_outputs(
            output_folder=output_folder,
            archive_folder=archive_folder,
        )
    )

    qualification_data = read_all_qualifications()
    dashboard_results = calculate_dashboard_results()

    timestamp = datetime.now().strftime(
        "%Y-%m-%d"
    )

    output_path = output_folder / (
        f"Dark Universe Qualification Coverage "
        f"{timestamp}.xlsx"
    )

    workbook = Workbook()

    create_dashboard(
        workbook,
        dashboard_results,
    )

    create_technician_sheet(
        workbook,
        qualification_data,
    )

    create_venue_summary(
        workbook,
        dashboard_results,
    )

    create_exceptions_sheet(
        workbook,
        qualification_data,
    )

    create_run_summary(
        workbook,
        qualification_data,
        dashboard_results,
        output_path,
    )

    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"

    workbook.save(output_path)

    if archived_count:
        print(
            "Archived previous workbooks: "
            f"{archived_count}"
        )
        print(f"Archive: {archive_path}")
    else:
        print("Archived previous workbooks: 0")

    exception_count = sum(
        len(exception_list)
        for exception_list
        in dashboard_results["exceptions"].values()
    )

    print("Qualification report created successfully.")
    print(f"Output: {output_path}")
    print(
        "Technicians: "
        f"{len(qualification_data['technicians'])}"
    )
    print(f"Exceptions: {exception_count}")


if __name__ == "__main__":
    generate_report()
if __name__ == "__main__":
    generate_report()