from pathlib import Path
import re

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parent.parent
LINE_SCHEDULE_FOLDER = PROJECT_ROOT / "Input" / "Line Schedule"

EXPECTED_SHEETS = {
    "1st Shift": "Operational - 1st Shift",
    "2nd Shift": "Operational - 2nd Shift",
    "3rd Shift Unchained": "3rd Shift Unchained",
    "3rd Shift Wolf": "3rd Shift Wolf",
}

EXCLUDED_LABELS = {
    "technician",
    "lead",
    "show controls",
    "mechanics",
    "nightly total",
    "daily total",
}

OPEN_POSITION_PREFIXES = (
    "open",
    "open -",
)


def find_line_schedule() -> Path:
    files = [
        file
        for file in LINE_SCHEDULE_FOLDER.glob("*.xlsx")
        if not file.name.startswith("~$")
    ]

    if not files:
        raise FileNotFoundError(
            f"No Line Schedule workbook found in: {LINE_SCHEDULE_FOLDER}"
        )

    if len(files) > 1:
        file_names = "\n".join(f"  - {file.name}" for file in files)
        raise RuntimeError(
            "More than one Line Schedule workbook was found:\n"
            f"{file_names}\n"
            "Keep only one current workbook in the folder."
        )

    return files[0]


def clean_name(value) -> str:
    if value is None:
        return ""

    name = str(value).strip()
    name = name.replace("*", "")
    name = re.sub(r"\s+", " ", name)

    return name


def clean_sso(value) -> str:
    if value is None:
        return ""

    if isinstance(value, int):
        sso = str(value)
    elif isinstance(value, float) and value.is_integer():
        sso = str(int(value))
    else:
        sso = str(value).strip()
        sso = re.sub(r"\.0$", "", sso)

    return sso if re.fullmatch(r"\d{9}", sso) else ""


def is_excluded_name(name: str) -> bool:
    normalized_name = name.casefold()

    if not normalized_name:
        return True

    if normalized_name in EXCLUDED_LABELS:
        return True

    return normalized_name.startswith(OPEN_POSITION_PREFIXES)


def read_line_schedule() -> tuple[list[dict], list[dict]]:
    workbook_path = find_line_schedule()

    workbook = load_workbook(
        workbook_path,
        read_only=True,
        data_only=True,
    )

    technicians = []
    exceptions = []

    for worksheet_name, report_group in EXPECTED_SHEETS.items():
        if worksheet_name not in workbook.sheetnames:
            exceptions.append(
                {
                    "sheet": worksheet_name,
                    "row": "",
                    "name": "",
                    "issue": "Expected worksheet is missing",
                }
            )
            continue

        worksheet = workbook[worksheet_name]

        for row_number, row in enumerate(
            worksheet.iter_rows(values_only=True),
            start=1,
        ):
            raw_name = row[0] if len(row) >= 1 else None
            raw_sso = row[3] if len(row) >= 4 else None

            name = clean_name(raw_name)
            sso = clean_sso(raw_sso)

            if is_excluded_name(name):
                continue

            if sso:
                technicians.append(
                    {
                        "name": name,
                        "sso": sso,
                        "schedule_group": report_group,
                        "source_sheet": worksheet_name,
                        "source_row": row_number,
                    }
                )
            elif row_number >= 4:
                exceptions.append(
                    {
                        "sheet": worksheet_name,
                        "row": row_number,
                        "name": name,
                        "issue": "Named line-schedule entry has no valid SSO",
                    }
                )

    workbook.close()

    return technicians, exceptions


def display_results() -> None:
    technicians, exceptions = read_line_schedule()

    print(f"Valid technicians: {len(technicians)}")
    print()

    for group in EXPECTED_SHEETS.values():
        group_technicians = [
            technician
            for technician in technicians
            if technician["schedule_group"] == group
        ]

        print(f"{group}: {len(group_technicians)}")

        for technician in group_technicians:
            print(
                f"  {technician['sso']} | "
                f"{technician['name']}"
            )

        print()

    print(f"Exceptions: {len(exceptions)}")

    for exception in exceptions:
        print(
            f"  {exception['sheet']} | "
            f"Row {exception['row']} | "
            f"{exception['name']} | "
            f"{exception['issue']}"
        )


if __name__ == "__main__":
    display_results()